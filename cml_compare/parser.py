"""Parser for Rolls-Royce MRO Characteristics Matrix (CM) workbooks.

The CM sheet layout (verified against template v01.x):

- Rows 1-8   : document metadata block (part name, engine, repair scheme, ...)
- Row ~10    : subtask descriptions (G2) for the section-G process grid
- Row ~14    : column header row; each header carries a reference code such
               as "Lower Tol (B2)" or "Subtask No. (G1)"
- Rows 15+   : one feature (balloon) per row

Sections A-F occupy fixed columns identified by their reference codes
(A1..A7, B1..B9, C1..C7, D1..D7, E1..E7, F1..F4). Section G is a
feature-vs-subtask grid whose columns differ between revisions, so grid
cells are keyed by subtask number rather than column position.

Nothing here hardcodes cell addresses: the header row and every column are
located by scanning for their reference codes, which tolerates template
drift between CM revisions.
"""

from __future__ import annotations

import datetime as _dt
import re
from dataclasses import dataclass, field

import openpyxl
from openpyxl.utils import get_column_letter

# Reference codes for the fixed (per-feature) columns of sections A-F.
FIXED_REF_CODES = (
    [f"A{i}" for i in range(1, 8)]
    + [f"B{i}" for i in range(1, 10)]
    + [f"C{i}" for i in range(1, 8)]
    + [f"D{i}" for i in range(1, 8)]
    + [f"E{i}" for i in range(1, 8)]
    + [f"F{i}" for i in range(1, 5)]
)

# Human-readable names for every reference code (used in reports).
FIELD_NAMES = {
    "A1": "Feature Type",
    "A2": "Engine Manual Ref / Task",
    "A3": "Subtask Number/s",
    "A4": "Feature No.",
    "A5": "No. of sub-features",
    "A6": "No. of sub-sub-features",
    "A7": "Fig No.",
    "B1": "Dim / other",
    "B2": "Lower Tol",
    "B3": "Upper Tol",
    "B4": "Control Band Dim",
    "B5": "Lower Control Limit",
    "B6": "Upper Control Limit",
    "B7": "Feature Description",
    "B8": "Feature Code",
    "B9": "Feature Notes",
    "C1": "Severity",
    "C2": "Occurrence",
    "C3": "Detection",
    "C4": "RPN",
    "C5": "S,L,F,P,W",
    "C6": "Inspection Category",
    "C7": "FVRA Notes",
    "D1": "Primary Datum",
    "D2": "Secondary Datum",
    "D3": "Tertiary Datum",
    "D4": "Thermally Sensitive?",
    "D5": "Thermal Offset",
    "D6": "Coefficient of Expansion",
    "D7": "dT for 10%*Tol Thermal Error",
    "E1": "Verification Method",
    "E2": "Equipment Type",
    "E3": "Equipment Identification",
    "E4": "MSA Verification Method",
    "E5": "MSA Report Number",
    "E6": "MSA Result",
    "E7": "Verification Notes",
    "F1": "Over check Method",
    "F2": "Over check Equipment Type",
    "F3": "Over check Equipment Id",
    "F4": "Over check Notes",
    "G4": "Final Feature Verification",
    "G5": "CM Notes",
}

_REF_CODE_RE = re.compile(r"\(\s*([A-G][1-9])\s*\)")
# Group-header columns in the G grid, e.g. "(4)    Part 1"
_GROUP_HEADER_RE = re.compile(r"^\(\d+\)")
# Subtask column headers, e.g. "(A) 72-41-13-350-387"
_SUBTASK_RE = re.compile(r"^\(([A-Z]{1,2})\)\s*(\S.*)$")


@dataclass
class Subtask:
    """One column of the section-G process grid."""

    key: str          # unique key: "<group> / <task number>"
    number: str       # task number, e.g. "72-41-13-350-387"
    group: str        # group header text, e.g. "(4) Part 1"
    description: str  # subtask description from the G2 row
    column: int       # 1-based column index in the source sheet


@dataclass
class Feature:
    """One feature (balloon) row of the CM sheet."""

    feature_no: str
    row: int                                  # 1-based row in the source sheet
    fields: dict = field(default_factory=dict)         # ref code -> raw value
    subtask_codes: dict = field(default_factory=dict)  # subtask key -> code (X/M/P/IC/...)


@dataclass
class CMLDocument:
    """Parsed content of one CM workbook."""

    path: str
    metadata: dict = field(default_factory=dict)       # label -> value
    version: str = ""                                  # latest CM version, e.g. "v03.00.00"
    features: dict = field(default_factory=dict)       # feature_no -> Feature
    subtasks: dict = field(default_factory=dict)       # subtask key -> Subtask
    # Layout info, reused by the report writer to locate cells in the new file
    sheet_name: str = "CM"
    header_row: int = 0
    columns: dict = field(default_factory=dict)        # ref code -> 1-based column
    first_data_row: int = 0
    last_data_row: int = 0

    def feature_order(self):
        return sorted(self.features.values(), key=lambda f: f.row)


class CMLParseError(Exception):
    pass


def _cell_text(value) -> str:
    if value is None:
        return ""
    return str(value)


def _find_cm_sheet(wb):
    for name in wb.sheetnames:
        if name.strip().upper() == "CM":
            return wb[name]
    # Fall back to the sheet containing a header cell with "(A4)"
    for name in wb.sheetnames:
        ws = wb[name]
        for row in ws.iter_rows(min_row=1, max_row=40):
            for cell in row:
                if "(A4)" in _cell_text(cell.value):
                    return ws
    raise CMLParseError("Could not find a 'CM' sheet in the workbook")


def _find_header_row(ws) -> int:
    """The header row is the one whose cells carry the (A4) reference code."""
    for row in ws.iter_rows(min_row=1, max_row=60):
        for cell in row:
            if "(A4)" in _cell_text(cell.value):
                return cell.row
    raise CMLParseError("Could not locate the CM header row (no '(A4)' marker found)")


def _map_columns(ws, header_row: int) -> dict:
    """Map reference codes (A1..F4, G1, G4, G5) to column indices.

    A1..F4 and G1/G5 live on the header row. G4 ("Final Feature Verification
    Method") is labelled a few rows above the header row, so those rows are
    scanned as well.
    """
    columns = {}
    for cell in ws[header_row]:
        for code in _REF_CODE_RE.findall(_cell_text(cell.value)):
            columns.setdefault(code, cell.column)
    for r in range(max(1, header_row - 8), header_row):
        for cell in ws[r]:
            for code in _REF_CODE_RE.findall(_cell_text(cell.value)):
                if code in ("G2", "G3", "G4"):
                    columns.setdefault(code, cell.column)
    missing = [c for c in ("A1", "A4", "B7") if c not in columns]
    if missing:
        raise CMLParseError(f"CM header row is missing reference codes: {missing}")
    return columns


def _parse_subtasks(ws, header_row: int, columns: dict) -> dict:
    """Read the section-G grid columns between G1 and G4."""
    g1_col = columns.get("G1")
    g4_col = columns.get("G4")
    if not g1_col or not g4_col:
        return {}
    # Subtask descriptions (G2) sit on the row that holds the "(G2)" label.
    desc_row = None
    for r in range(max(1, header_row - 8), header_row):
        for cell in ws[r]:
            if "(G2)" in _cell_text(cell.value):
                desc_row = r
                break
        if desc_row:
            break

    subtasks = {}
    group = ""
    for col in range(g1_col + 1, g4_col):
        header = _cell_text(ws.cell(header_row, col).value).strip()
        if not header:
            continue
        m = _SUBTASK_RE.match(header)
        if not m and _GROUP_HEADER_RE.match(header):
            group = re.sub(r"\s+", " ", header)
            continue
        number = re.sub(r"\s+", " ", m.group(2).strip()) if m else re.sub(r"\s+", " ", header)
        key = f"{group} / {number}" if group else number
        if key in subtasks:  # same task repeated within a group: disambiguate
            n = 2
            while f"{key} #{n}" in subtasks:
                n += 1
            key = f"{key} #{n}"
        description = ""
        if desc_row:
            description = re.sub(r"\s+", " ", _cell_text(ws.cell(desc_row, col).value)).strip()
        subtasks[key] = Subtask(key=key, number=number, group=group,
                                description=description, column=col)
    return subtasks


def _feature_key(value) -> str:
    """Normalise a Feature No. cell into a stable dictionary key."""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _metadata_from_block(ws, max_row: int) -> dict:
    """Read "label: value" rows from a metadata block.

    The value is taken from the first populated cell within 4 columns of its
    label; the cap keeps the scan away from unrelated content further right
    (e.g. the feature-type legend on the CM sheet).
    """
    metadata = {}
    for row in ws.iter_rows(min_row=1, max_row=max_row):
        cells = [c for c in row if c.value is not None and _cell_text(c.value).strip()]
        for i, c in enumerate(cells):
            label = _cell_text(c.value).strip()
            is_secondary = label.rstrip(":").strip().lower() in ("issue / rev date", "vendor")
            if i > 0 and not is_secondary:
                continue
            label = label.rstrip(":").strip()
            for nxt in cells[i + 1:]:
                if nxt.column > c.column + 4:
                    break
                v = nxt.value
                text = _cell_text(v).strip()
                if text.endswith(":"):
                    break  # ran into the next label on this row
                if text and text != "0":
                    if isinstance(v, _dt.datetime):
                        text = v.date().isoformat()
                    metadata[label] = text
                    break
    return metadata


def _parse_metadata(wb, ws, header_row: int) -> tuple[dict, str]:
    """Collect document metadata plus the latest CM version.

    The CM sheet's metadata cells are formulas (defined names pointing at
    the 'CM Version Control' sheet), so that sheet's literal values are
    taken as the base and any cached values from the CM sheet overlaid.
    """
    metadata = {}
    version = ""
    for name in wb.sheetnames:
        if "version control" not in name.lower() or "template" in name.lower():
            continue
        vc = wb[name]
        version_start = vc.max_row
        for row in vc.iter_rows():
            text = _cell_text(row[0].value).strip()
            if text.lower().startswith("version control"):
                version_start = row[0].row
            if re.match(r"^v\d+\.\d+", text, re.IGNORECASE):
                version = text  # keep the last (latest) entry
        metadata.update(_metadata_from_block(vc, version_start))
        break
    metadata.update(_metadata_from_block(ws, max(1, header_row - 5)))
    return metadata, version


def parse_cml(path: str) -> CMLDocument:
    """Parse a CML workbook into a CMLDocument (values only, read-only)."""
    wb = openpyxl.load_workbook(path, data_only=True, read_only=False)
    try:
        ws = _find_cm_sheet(wb)
        header_row = _find_header_row(ws)
        columns = _map_columns(ws, header_row)
        subtasks = _parse_subtasks(ws, header_row, columns)
        metadata, version = _parse_metadata(wb, ws, header_row)

        doc = CMLDocument(
            path=path,
            metadata=metadata,
            version=version,
            subtasks=subtasks,
            sheet_name=ws.title,
            header_row=header_row,
            columns=dict(columns),
            first_data_row=header_row + 1,
        )

        feat_col = columns["A4"]
        type_col = columns["A1"]
        for r in range(header_row + 1, ws.max_row + 1):
            feat_val = ws.cell(r, feat_col).value
            type_val = ws.cell(r, type_col).value
            if feat_val is None and type_val is None:
                continue
            key = _feature_key(feat_val) if feat_val is not None else f"row{r}"
            if key in doc.features:  # duplicate balloon number: keep both, flagged
                key = f"{key} (row {r})"
            feature = Feature(feature_no=key, row=r)
            for code in FIXED_REF_CODES + ["G4", "G5"]:
                col = columns.get(code)
                if col:
                    feature.fields[code] = ws.cell(r, col).value
            for st in subtasks.values():
                v = ws.cell(r, st.column).value
                if v is not None and _cell_text(v).strip():
                    feature.subtask_codes[st.key] = _cell_text(v).strip()
            doc.features[key] = feature
            doc.last_data_row = r
        return doc
    finally:
        wb.close()


def column_letter(index: int) -> str:
    return get_column_letter(index)
