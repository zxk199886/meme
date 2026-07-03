"""Report writer: marked-up copy of the NEW CML with differences highlighted.

The new workbook is copied as-is (formulas and formatting preserved), then:

- every changed cell on the CM sheet gets a colour fill (by change
  category) and a cell comment stating the old value,
- rows of newly added features are filled green,
- a "Comparison Summary" sheet is prepended listing metadata changes,
  added/removed features (removed rows no longer exist in the new file, so
  their key data is tabulated there), every field change, and the legend.
"""

from __future__ import annotations

import datetime as _dt

import openpyxl
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from .diff import (
    CATEGORY_DATUM,
    CATEGORY_DIMENSION,
    CATEGORY_FVRA,
    CATEGORY_PROCESS,
    CATEGORY_TEXT,
    CATEGORY_TOLERANCE,
    CATEGORY_TOOLING,
    DiffResult,
    display_value,
)

SUMMARY_SHEET = "Comparison Summary"
COMMENT_AUTHOR = "CML Compare"

CATEGORY_COLORS = {
    CATEGORY_DIMENSION: "FFC000",  # orange
    CATEGORY_TOLERANCE: "FF7043",  # deep orange
    CATEGORY_TOOLING: "9DC3E6",    # blue
    CATEGORY_FVRA: "D9A7E0",       # purple
    CATEGORY_DATUM: "BFBFBF",      # grey
    CATEGORY_PROCESS: "F8B4D9",    # pink
    CATEGORY_TEXT: "FFE699",       # yellow
}
ADDED_COLOR = "A9D08E"    # green
REMOVED_COLOR = "FF9999"  # red (summary sheet only)


def _fill(color: str) -> PatternFill:
    return PatternFill(fill_type="solid", start_color=color, end_color=color)


def _comment(text: str) -> Comment:
    c = Comment(text, COMMENT_AUTHOR)
    c.width, c.height = 260, 90
    return c


def write_marked_up_report(diff: DiffResult, output_path: str) -> str:
    """Write the marked-up copy of the new CML. Returns output_path."""
    wb = openpyxl.load_workbook(diff.new_doc.path)  # formulas + formatting kept
    ws = wb[diff.new_doc.sheet_name]
    new_doc, old_doc = diff.new_doc, diff.old_doc
    old_label = old_doc.version or "previous revision"

    # --- highlight changed cells -------------------------------------------
    for change in diff.field_changes:
        feature = new_doc.features.get(change.feature_no)
        if feature is None:
            continue
        if change.ref_code.startswith("G3:"):
            st = new_doc.subtasks.get(change.ref_code[3:])
            col = st.column if st else None
        else:
            col = new_doc.columns.get(change.ref_code)
        if not col:
            continue
        cell = ws.cell(feature.row, col)
        cell.fill = _fill(CATEGORY_COLORS.get(change.category, CATEGORY_COLORS[CATEGORY_TEXT]))
        cell.comment = _comment(
            f"[{change.category}] {change.field_name}\n"
            f"Was ({old_label}): {change.old}\nNow: {change.new}")

    # --- highlight added feature rows ---------------------------------------
    max_col = max([new_doc.columns.get("G5") or 0, new_doc.columns.get("G4") or 0,
                   *(st.column for st in new_doc.subtasks.values())] or [50])
    for feature_no in diff.added_features:
        feature = new_doc.features[feature_no]
        for col in range(1, max_col + 1):
            ws.cell(feature.row, col).fill = _fill(ADDED_COLOR)
        ws.cell(feature.row, new_doc.columns["A4"]).comment = _comment(
            f"NEW FEATURE\nNot present in {old_label}.")

    # --- flag added subtask columns -----------------------------------------
    for st in diff.added_subtasks:
        cell = ws.cell(new_doc.header_row, st.column)
        cell.fill = _fill(ADDED_COLOR)
        cell.comment = _comment(f"NEW SUBTASK COLUMN\nNot present in {old_label}.")

    _write_summary_sheet(wb, diff)
    wb.save(output_path)
    return output_path


# --------------------------------------------------------------------------
# Summary sheet
# --------------------------------------------------------------------------

_TITLE_FONT = Font(bold=True, size=14)
_H_FONT = Font(bold=True, size=11)
_HDR_FONT = Font(bold=True, color="FFFFFF")
_HDR_FILL = _fill("305496")
_BORDER = Border(*(Side(style="thin", color="BFBFBF"),) * 4)


def _write_summary_sheet(wb, diff: DiffResult) -> None:
    if SUMMARY_SHEET in wb.sheetnames:
        del wb[SUMMARY_SHEET]
    ws = wb.create_sheet(SUMMARY_SHEET, 0)
    ws.sheet_view.showGridLines = False
    for col, width in zip("ABCDEFG", (16, 34, 44, 44, 22, 18, 18)):
        ws.column_dimensions[col].width = width

    old_doc, new_doc = diff.old_doc, diff.new_doc
    row = 1

    def put(r, c, value, font=None, fill=None, wrap=False):
        cell = ws.cell(r, c, value)
        if font:
            cell.font = font
        if fill:
            cell.fill = fill
        cell.alignment = Alignment(vertical="top", wrap_text=wrap)
        cell.border = _BORDER
        return cell

    def heading(text):
        nonlocal row
        row += 1
        ws.cell(row, 1, text).font = _H_FONT
        row += 1

    def table_header(*labels):
        nonlocal row
        for i, label in enumerate(labels, 1):
            put(row, i, label, font=_HDR_FONT, fill=_HDR_FILL)
        row += 1

    ws.cell(row, 1, "Characteristics Matrix - Revision Comparison").font = _TITLE_FONT
    row += 2
    meta_pairs = [
        ("Old revision", f"{old_doc.version or 'n/a'}  -  {_basename(old_doc.path)}"),
        ("New revision", f"{new_doc.version or 'n/a'}  -  {_basename(new_doc.path)}"),
        ("Part", new_doc.metadata.get("Commodity - Part Name", "")),
        ("Repair scheme", new_doc.metadata.get("Repair Scheme", "")),
        ("Compared on", _dt.date.today().isoformat()),
    ]
    for label, value in meta_pairs:
        put(row, 1, label, font=_H_FONT)
        put(row, 2, value)
        row += 1
    row += 1
    note = ("NOTE: This marked-up copy is a review aid, not a controlled document. "
            "Embedded images / OLE objects (e.g. balloon diagram attachments) are "
            "not carried over - refer to the controlled CML for those.")
    put(row, 1, note, font=Font(italic=True, color="7F7F7F"))
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=4)
    row += 1

    heading("Result Overview")
    common = set(old_doc.features) & set(new_doc.features)
    overview = [
        ("Features in old / new", f"{len(old_doc.features)} / {len(new_doc.features)}"),
        ("Added features", len(diff.added_features)),
        ("Removed features", len(diff.removed_features)),
        ("Modified features", len(diff.modified_features)),
        ("Unchanged features", len(common) - len(diff.modified_features)),
        ("Individual cell changes", len(diff.field_changes)),
        ("Metadata changes", len(diff.metadata_changes)),
        ("Subtask columns added / removed",
         f"{len(diff.added_subtasks)} / {len(diff.removed_subtasks)}"),
    ]
    for label, value in overview:
        put(row, 1, label, font=_H_FONT)
        put(row, 2, value)
        row += 1
    if not diff.has_changes:
        put(row, 1, "NO DIFFERENCES FOUND", font=Font(bold=True, color="006100"),
            fill=_fill("C6EFCE"))
        row += 1

    heading("Colour Legend (as used on the CM sheet)")
    legend = [("Added feature / subtask (whole row green)", ADDED_COLOR),
              ("Removed feature (listed below only)", REMOVED_COLOR)]
    legend += [(f"Changed cell - {cat}", color) for cat, color in CATEGORY_COLORS.items()]
    for label, color in legend:
        put(row, 1, "", fill=_fill(color))
        put(row, 2, label)
        row += 1

    if diff.metadata_changes:
        heading("Document / Metadata Changes")
        table_header("Field", "Old", "New")
        for ch in diff.metadata_changes:
            put(row, 1, ch.label)
            put(row, 2, ch.old, wrap=True)
            put(row, 3, ch.new, wrap=True)
            row += 1

    if diff.added_features:
        heading("Added Features (rows highlighted green on the CM sheet)")
        table_header("Feature No.", "Description", "Dim", "Lower / Upper Tol",
                     "Verification", "Equipment")
        for no in diff.added_features:
            f = new_doc.features[no]
            _feature_row(put, row, no, f, _fill(ADDED_COLOR))
            row += 1

    if diff.removed_features:
        heading("Removed Features (present in old revision only)")
        table_header("Feature No.", "Description", "Dim", "Lower / Upper Tol",
                     "Verification", "Equipment")
        for no in diff.removed_features:
            f = old_doc.features[no]
            _feature_row(put, row, no, f, _fill(REMOVED_COLOR))
            row += 1

    if diff.added_subtasks or diff.removed_subtasks:
        heading("Process Grid Subtask Columns Added / Removed")
        table_header("Change", "Subtask No.", "Description", "Group")
        for st in diff.added_subtasks:
            put(row, 1, "ADDED", fill=_fill(ADDED_COLOR))
            put(row, 2, st.number)
            put(row, 3, st.description, wrap=True)
            put(row, 4, st.group)
            row += 1
        for st in diff.removed_subtasks:
            put(row, 1, "REMOVED", fill=_fill(REMOVED_COLOR))
            put(row, 2, st.number)
            put(row, 3, st.description, wrap=True)
            put(row, 4, st.group)
            row += 1

    if diff.field_changes:
        heading("All Cell Changes")
        table_header("Feature No.", "Field", "Old", "New", "Category", "CM Cell")
        for ch in diff.field_changes:
            f = new_doc.features.get(ch.feature_no)
            if ch.ref_code.startswith("G3:"):
                st = new_doc.subtasks.get(ch.ref_code[3:])
                col = st.column if st else None
            else:
                col = new_doc.columns.get(ch.ref_code)
            ref = f"{get_column_letter(col)}{f.row}" if (f and col) else ""
            put(row, 1, ch.feature_no)
            put(row, 2, ch.field_name, wrap=True)
            put(row, 3, ch.old, wrap=True)
            put(row, 4, ch.new, wrap=True)
            put(row, 5, ch.category,
                fill=_fill(CATEGORY_COLORS.get(ch.category, CATEGORY_COLORS[CATEGORY_TEXT])))
            put(row, 6, ref)
            row += 1

    ws.freeze_panes = "A2"


def _feature_row(put, row, feature_no, feature, fill):
    g = feature.fields.get
    put(row, 1, feature_no, fill=fill)
    put(row, 2, display_value(g("B7")), wrap=True)
    put(row, 3, display_value(g("B1")))
    put(row, 4, f"{display_value(g('B2'))} / {display_value(g('B3'))}")
    put(row, 5, display_value(g("E1")))
    put(row, 6, display_value(g("E2")), wrap=True)


def _basename(path: str) -> str:
    import os
    return os.path.basename(path)
