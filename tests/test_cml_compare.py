"""Tests for the CML comparison tool.

Fixtures are small synthetic workbooks that mimic the Rolls-Royce CM
template layout (metadata block, G2 description row, reference-coded
header row, feature rows, section-G subtask grid). No proprietary data.
"""

import os
import sys
import tempfile
import unittest

import openpyxl

sys.path.insert(0, os.path.join(os.path.dirname(__file__), ".."))

from cml_compare.diff import CATEGORY_PROCESS, CATEGORY_TOLERANCE, CATEGORY_TOOLING, compare
from cml_compare.parser import parse_cml
from cml_compare.report import SUMMARY_SHEET, write_marked_up_report

HEADER_ROW = 14
DESC_ROW = 10

# (ref code, header label) for the fixed columns used in the fixture
FIXED_HEADERS = [
    ("A1", "Feature Type\n(A1)"), ("A2", "Engine Manual Ref / Task\n(A2)"),
    ("A3", "Subtask Number/s\n(A3)"), ("A4", "Feature No.\n(A4)"),
    ("A5", "No. of sub-features\n(A5)"), ("A6", "No. of sub-sub-features\n(A6)"),
    ("A7", "Fig No.\n(A7)"),
    ("B1", "Dim / other\n(B1)"), ("B2", "Lower Tol \n(B2)"), ("B3", "Upper Tol \n(B3)"),
    ("B4", "Dim / other \n(B4)"), ("B5", "Lower Control Limit \n(B5)"),
    ("B6", "Upper Control Limit\n(B6)"), ("B7", "Feature Desc.\n(B7)"),
    ("B8", "Feature Code\n(B8)"), ("B9", "Feature notes if required\n(B9)"),
    ("C1", "Severity\n(C1)"), ("C2", "Occurrence\n(C2)"), ("C3", "Detection\n(C3)"),
    ("C4", "RPN\n(C4)"), ("C5", "S,L,F,P,W\n(C5)"), ("C6", "Inspection Category \n(C6)"),
    ("C7", "FVRA Notes if required\n(C7)"),
    ("D1", "Primary (D1)"), ("D2", "Secondary (D2)"), ("D3", "Tertiary (D3)"),
    ("D4", "Thermally Sensitive Feature?\n(D4)"), ("D5", "Thermal Offset\n(D5)"),
    ("D6", "Coefficient of Expansion\n(D6)"), ("D7", "dT for 10%*Tol\n(D7)"),
    ("E1", "Verification Method\n(E1)"), ("E2", "Equipment Type\n(E2) "),
    ("E3", "Equipment Identification\n(E3)"), ("E4", "MSA Method\n(E4)"),
    ("E5", "MSA Report Number\n(E5)"), ("E6", "MSA Result\n(E6)"),
    ("E7", "Notes if required\n(E7)"),
    ("F1", "Over check Method\n(F1)"), ("F2", "Equipment Type\n(F2)"),
    ("F3", "Equipment Identification\n(F3)"), ("F4", "Notes if required\n(F4)"),
]

# Default per-feature values; tests override individual fields.
DEFAULT_FEATURE = {
    "A1": "E", "A2": "FRSK999_12-34-56", "A3": "12-34-56-350-001", "A5": "N/A",
    "A6": "N/A", "A7": "12-34-56-990-001", "B1": 2.04, "B2": 0, "B3": 0.12,
    "B7": "DIAMETER", "B8": "DIAM", "C1": 5, "C2": 1, "C3": 1, "C4": 5, "C6": "C",
    "D4": "YES", "E1": "MANUAL", "E2": "PLAIN PLUG GAUGE", "E4": "READ-ACROSS",
    "F1": "MANUAL", "F2": "PLAIN PLUG GAUGE", "G4": "IM",
}

BASE_SUBTASKS = [
    ("(4)    Part 1", None),
    ("(A) 12-34-56-350-001", "Weld the part"),
    ("(B) 12-34-56-220-002", "Examine the part"),
    ("(C) 12-34-56-230-003", "Penetrant inspection"),
]

BASE_METADATA = {
    "Commodity - Part Name": "TEST LINER ASSEMBLY",
    "Description": "REPAIR WALL DEFECTS",
    "Engine Type (s)": "TRENT 900",
    "Engine Manual": "CIR 12-34-56-300-999",
    "Repair Scheme": "FRSK999",
    "Part Number(s)": "FW11111, FW22222",
}


def build_cml(path, features, subtasks=BASE_SUBTASKS, metadata=None, version="v01.00.00"):
    """Write a synthetic CM workbook.

    features: dict feature_no -> dict of overrides; may include
              "grid": {subtask header -> code}.
    """
    metadata = {**BASE_METADATA, **(metadata or {})}
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "CM"
    for r, (label, value) in enumerate(metadata.items(), 2):
        ws.cell(r, 1, label)
        ws.cell(r, 4, value)

    col_of = {}
    col = 1
    for code, header in FIXED_HEADERS:
        ws.cell(HEADER_ROW, col, header)
        col_of[code] = col
        col += 1
    ws.cell(HEADER_ROW, col, "Subtask No. (G1)")
    ws.cell(DESC_ROW, col, "Subtask Description (G2)")
    col += 1
    subtask_cols = {}
    for header, desc in subtasks:
        ws.cell(HEADER_ROW, col, header)
        if desc:
            ws.cell(DESC_ROW, col, desc)
        subtask_cols[header] = col
        col += 1
    ws.cell(DESC_ROW, col, "Final Feature Verification Method\n(G4)")
    col_of["G4"] = col
    col += 1
    ws.cell(HEADER_ROW, col, "Notes if required\n(G5)")
    col_of["G5"] = col

    for i, (feature_no, overrides) in enumerate(features.items()):
        r = HEADER_ROW + 1 + i
        values = {**DEFAULT_FEATURE, **{k: v for k, v in overrides.items() if k != "grid"}}
        ws.cell(r, col_of["A4"], feature_no)
        for code, value in values.items():
            if code in col_of:
                ws.cell(r, col_of[code], value)
        for header, code in overrides.get("grid", {}).items():
            ws.cell(r, subtask_cols[header], code)

    vc = wb.create_sheet("CM Version Control")
    vc["A1"] = "Version Control"
    vc["A2"] = version
    vc["B2"] = "2026-01-01"
    wb.save(path)


BASE_FEATURES = {
    1: {"grid": {"(A) 12-34-56-350-001": "X", "(B) 12-34-56-220-002": "IM"}},
    2: {"B1": 5.5, "B7": "MINIMUM LENGTH", "B8": "LENG", "B2": "N/A", "B3": "N/A",
        "E2": "CALIPER - DIGITAL",
        "grid": {"(B) 12-34-56-220-002": "IM"}},
    3: {"B1": 0.25, "B7": "PROFILE OF A SURFACE", "B8": "PROS", "C2": 5, "C3": 5,
        "C4": 125, "C6": "B", "E1": "CMS", "E2": "CMM",
        "grid": {"(C) 12-34-56-230-003": "IC"}},
    4: {"B7": "REMOVE SHARP EDGES", "B8": "OTHR", "E1": "VISUAL",
        "E2": "VISUAL VERIFICATION", "G4": "IV"},
}


class CMLCompareTest(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.tmp = tempfile.TemporaryDirectory(prefix="cml_test_")
        d = cls.tmp.name
        cls.old_path = os.path.join(d, "old.xlsx")
        cls.new_path = os.path.join(d, "new.xlsx")

        build_cml(cls.old_path, BASE_FEATURES, version="v01.00.00")

        # New revision with known edits:
        new_features = {k: dict(v) for k, v in BASE_FEATURES.items()}
        for k in new_features:
            new_features[k]["grid"] = dict(BASE_FEATURES[k].get("grid", {}))
        new_features[1]["B3"] = 0.15                        # tolerance change
        new_features[2]["E2"] = "MICROMETER - DIGITAL"      # equipment change
        new_features[3]["grid"]["(C) 12-34-56-230-003"] = "IM"  # grid code change
        del new_features[4]                                 # removed feature
        new_features[5] = {"B1": 3.04, "B7": "DIAMETER"}    # added feature
        new_subtasks = BASE_SUBTASKS + [("(D) 12-34-56-110-004", "Clean the part")]
        build_cml(cls.new_path, new_features, subtasks=new_subtasks,
                  metadata={"Part Number(s)": "FW11111, FW22222, FW33333"},
                  version="v02.00.00")

        cls.old_doc = parse_cml(cls.old_path)
        cls.new_doc = parse_cml(cls.new_path)
        cls.diff = compare(cls.old_doc, cls.new_doc)

    @classmethod
    def tearDownClass(cls):
        cls.tmp.cleanup()

    # --- parsing ----------------------------------------------------------
    def test_parse_features_and_metadata(self):
        self.assertEqual(set(self.old_doc.features), {"1", "2", "3", "4"})
        self.assertEqual(self.old_doc.version, "v01.00.00")
        self.assertEqual(self.old_doc.metadata["Repair Scheme"], "FRSK999")
        f1 = self.old_doc.features["1"]
        self.assertEqual(f1.fields["B1"], 2.04)
        self.assertEqual(f1.fields["B3"], 0.12)
        self.assertEqual(f1.fields["E2"], "PLAIN PLUG GAUGE")

    def test_parse_subtask_grid(self):
        self.assertEqual(len(self.old_doc.subtasks), 3)  # group header skipped
        keys = list(self.old_doc.subtasks)
        self.assertTrue(all(k.startswith("(4) Part 1 / ") for k in keys))
        f1 = self.old_doc.features["1"]
        self.assertEqual(set(f1.subtask_codes.values()), {"X", "IM"})

    # --- diffing ----------------------------------------------------------
    def test_added_and_removed_features(self):
        self.assertEqual(self.diff.added_features, ["5"])
        self.assertEqual(self.diff.removed_features, ["4"])

    def test_tolerance_change_detected(self):
        changes = [c for c in self.diff.field_changes
                   if c.feature_no == "1" and c.ref_code == "B3"]
        self.assertEqual(len(changes), 1)
        self.assertEqual(changes[0].category, CATEGORY_TOLERANCE)
        self.assertEqual((changes[0].old, changes[0].new), ("0.12", "0.15"))

    def test_equipment_change_detected(self):
        changes = [c for c in self.diff.field_changes
                   if c.feature_no == "2" and c.ref_code == "E2"]
        self.assertEqual(len(changes), 1)
        self.assertEqual(changes[0].category, CATEGORY_TOOLING)
        self.assertEqual(changes[0].new, "MICROMETER - DIGITAL")

    def test_grid_change_detected(self):
        changes = [c for c in self.diff.field_changes
                   if c.feature_no == "3" and c.ref_code.startswith("G3:")]
        self.assertEqual(len(changes), 1)
        self.assertEqual(changes[0].category, CATEGORY_PROCESS)
        self.assertEqual((changes[0].old, changes[0].new), ("IC", "IM"))

    def test_subtask_column_added(self):
        self.assertEqual([st.number for st in self.diff.added_subtasks],
                         ["12-34-56-110-004"])
        self.assertEqual(self.diff.removed_subtasks, [])

    def test_metadata_change_detected(self):
        labels = {m.label: (m.old, m.new) for m in self.diff.metadata_changes}
        self.assertIn("Part Number(s)", labels)
        self.assertEqual(labels["CM Version"], ("v01.00.00", "v02.00.00"))

    def test_no_false_positives(self):
        expected = {("1", "B3"), ("2", "E2")}
        got = {(c.feature_no, c.ref_code) for c in self.diff.field_changes
               if not c.ref_code.startswith("G3:")}
        self.assertEqual(got, expected)

    def test_identical_files_have_no_changes(self):
        diff = compare(parse_cml(self.old_path), parse_cml(self.old_path))
        self.assertFalse(diff.has_changes)

    def test_numeric_and_whitespace_normalisation(self):
        d = self.tmp.name
        a, b = os.path.join(d, "a.xlsx"), os.path.join(d, "b.xlsx")
        build_cml(a, {1: {"B3": 0.4, "B7": "TRUE POSITION\nDIAMETER"}})
        build_cml(b, {1: {"B3": "0.40", "B7": "TRUE POSITION DIAMETER"}})
        self.assertFalse(compare(parse_cml(a), parse_cml(b)).has_changes)

    # --- report -----------------------------------------------------------
    def test_marked_up_report(self):
        out = os.path.join(self.tmp.name, "report.xlsx")
        write_marked_up_report(self.diff, out)
        wb = openpyxl.load_workbook(out)
        self.assertEqual(wb.sheetnames[0], SUMMARY_SHEET)
        ws = wb["CM"]
        # Changed tolerance cell of feature 1 must be filled and commented
        f1 = self.new_doc.features["1"]
        cell = ws.cell(f1.row, self.new_doc.columns["B3"])
        self.assertEqual(cell.fill.fill_type, "solid")
        self.assertIsNotNone(cell.comment)
        self.assertIn("0.12", cell.comment.text)
        # Added feature row is filled
        f5 = self.new_doc.features["5"]
        self.assertEqual(ws.cell(f5.row, 1).fill.fill_type, "solid")
        # Removed feature appears on the summary sheet
        summary_text = "\n".join(str(c.value) for row in wb[SUMMARY_SHEET].iter_rows()
                                 for c in row if c.value is not None)
        self.assertIn("Removed Features", summary_text)
        self.assertIn("REMOVE SHARP EDGES", summary_text)


if __name__ == "__main__":
    unittest.main(verbosity=2)
